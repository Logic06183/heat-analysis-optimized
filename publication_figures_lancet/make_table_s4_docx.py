"""Build a paste-ready Word table for Table S4 from the exported CSV."""
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches

HERE = Path(__file__).resolve().parent
df = pd.read_csv(HERE / "table_s4_interaction_pairs.csv")


def fmt_lag(s: str) -> str:
    s = str(s)
    return f"{s[:-1]} d" if s.endswith("d") and s[:-1].isdigit() else s


disp = pd.DataFrame({
    "#": df["rank"].astype(int),
    "Biomarker": df["biomarker"],
    "Effect modifier": df["modifier"],
    "Category": df["modifier_category"],
    "Temp lag": df["temperature_feature"].map(fmt_lag),
    "|SHAP interaction|": df["mean_abs_interaction"].map(lambda v: f"{v:.4f}"),
    "BH q": df["q_value_BH"].map(lambda v: f"{v:.3f}"),
})

WIDTHS = {  # inches; sum ~6.3" fits a portrait page
    "#": 0.35, "Biomarker": 1.05, "Effect modifier": 1.75, "Category": 0.85,
    "Temp lag": 0.6, "|SHAP interaction|": 1.05, "BH q": 0.6,
}
RIGHT = {"#", "|SHAP interaction|", "BH q"}
CENTER = {"Category", "Temp lag"}

doc = Document()
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(9)

title = doc.add_paragraph()
tr = title.add_run(
    "Table S4. FDR-significant temperature × effect-modifier interactions "
    "(Stage 2, ERA5-Land)."
)
tr.bold = True
tr.font.size = Pt(10)

cols = list(disp.columns)
table = doc.add_table(rows=len(disp) + 1, cols=len(cols))
table.style = "Table Grid"
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False

# Header row
for j, col in enumerate(cols):
    cell = table.rows[0].cells[j]
    cell.text = ""
    run = cell.paragraphs[0].add_run(col)
    run.bold = True
    run.font.size = Pt(9)
    cell.paragraphs[0].alignment = (
        WD_ALIGN_PARAGRAPH.RIGHT if col in RIGHT
        else WD_ALIGN_PARAGRAPH.CENTER if col in CENTER
        else WD_ALIGN_PARAGRAPH.LEFT
    )
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "D9E1F2")
    cell._tc.get_or_add_tcPr().append(shd)

# Data rows
for i in range(len(disp)):
    for j, col in enumerate(cols):
        cell = table.rows[i + 1].cells[j]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(str(disp.iloc[i, j]))
        run.font.size = Pt(8.5)
        p.alignment = (
            WD_ALIGN_PARAGRAPH.RIGHT if col in RIGHT
            else WD_ALIGN_PARAGRAPH.CENTER if col in CENTER
            else WD_ALIGN_PARAGRAPH.LEFT
        )

# Fixed column widths (must set on every cell + disable autofit)
for j, col in enumerate(cols):
    w = Inches(WIDTHS[col])
    for row in table.rows:
        row.cells[j].width = w

foot = doc.add_paragraph()
fr = foot.add_run(
    "|SHAP interaction|, mean absolute SHAP interaction value (top-50 candidate "
    "pairs per biomarker, ranked across the bootstrap models); BH q, "
    "Benjamini–Hochberg adjusted p-value, significant at q ≤ 0.10; Temp "
    "lag, temperature lag (days) at which the interaction was evaluated. "
    "n = 32 significant interaction pairs across 8 biomarkers."
)
fr.italic = True
fr.font.size = Pt(8)

out = HERE / "Table_S4_interaction_pairs.docx"
doc.save(out)
print(f"Saved: {out}  ({len(disp)} rows)")
